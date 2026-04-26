import os

MAX_PAGES = 10
MAX_CHARS = 5000


def extract_pdf(file_path: str) -> str:
    try:
        import pdfplumber
        texts = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages[:MAX_PAGES]:
                try:
                    text = page.extract_text()
                    if text:
                        texts.append(text)
                finally:
                    del page
        result = "\n\n".join(texts)
        return result[:MAX_CHARS] if len(result) > MAX_CHARS else result
    except Exception:
        return ""


def extract_docx(file_path: str) -> str:
    try:
        from docx import Document
        doc = Document(file_path)
        result = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return result[:MAX_CHARS] if len(result) > MAX_CHARS else result
    except Exception:
        return ""


def extract_pptx(file_path: str) -> str:
    try:
        from pptx import Presentation
        prs = Presentation(file_path)
        texts = []
        for i, slide in enumerate(prs.slides, 1):
            slide_texts = [f"[Slide {i}]"]
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for para in shape.text_frame.paragraphs:
                        line = para.text.strip()
                        if line:
                            slide_texts.append(line)
            if len(slide_texts) > 1:
                texts.append("\n".join(slide_texts))
        result = "\n\n".join(texts)
        return result[:MAX_CHARS] if len(result) > MAX_CHARS else result
    except Exception:
        return ""


def extract_text(raw_text: str) -> str:
    raw_text = raw_text.strip()
    return raw_text[:MAX_CHARS] if len(raw_text) > MAX_CHARS else raw_text


def extract_from_file(file_path: str, filename: str) -> str:
    ext = os.path.splitext(filename)[1].lower()
    if ext == ".pdf":
        return extract_pdf(file_path)
    elif ext == ".docx":
        return extract_docx(file_path)
    elif ext == ".pptx":
        return extract_pptx(file_path)
    elif ext == ".txt":
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                result = f.read().strip()
            return result[:MAX_CHARS] if len(result) > MAX_CHARS else result
        except Exception:
            return ""
    else:
        raise ValueError(f"Unsupported file type: {ext}")
